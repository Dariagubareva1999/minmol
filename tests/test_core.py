import tempfile
import unittest
import os
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from minmol.core import (
    calculate_target,
    classify_event,
    build_press_release,
    default_data,
    event_is_complete,
    normalize_data,
    parse_event_date,
    reporting_period,
    reporting_period_bounds,
)
from minmol.reports import export_all, export_appendix_2, export_press_release, preview_appendix_2
from minmol.security import delete_secret, load_secret, save_secret


class CoreTests(unittest.TestCase):
    @unittest.skipUnless(os.name == "nt", "Windows DPAPI is required")
    def test_api_key_is_encrypted_for_current_windows_user(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "key.bin"
            save_secret("new-secret-key", path)
            self.assertNotIn(b"new-secret-key", path.read_bytes())
            self.assertEqual(load_secret(path), "new-secret-key")
            delete_secret(path)
            self.assertFalse(path.exists())

    def test_parse_date(self):
        self.assertEqual(parse_event_date("15.03.2026").isoformat(), "2026-03-15")
        self.assertEqual(parse_event_date("2026-03-15").isoformat(), "2026-03-15")
        self.assertIsNone(parse_event_date("31.02.2026"))

    def test_reporting_period_changes_on_25th(self):
        self.assertEqual(reporting_period(parse_event_date("24.01.2026")), (2026, 0))
        self.assertEqual(reporting_period(parse_event_date("25.01.2026")), (2026, 1))
        self.assertEqual(reporting_period(parse_event_date("25.12.2026")), (2027, 0))
        self.assertEqual(
            tuple(value.isoformat() for value in reporting_period_bounds(2026, 1)),
            ("2026-01-25", "2026-02-24"),
        )

    def test_calculation_with_carry(self):
        data = default_data(2026)
        data["targets"][0]["plan"][:3] = [100, 250, 300]
        data["events"] = [
            {"date": "10.01.2026", "target_ids": ["01"], "participants": 60},
            {"date": "10.02.2026", "target_ids": ["01"], "participants": 180},
        ]
        rows = calculate_target(data, "01")
        self.assertEqual((rows[0]["monthly_plan"], rows[0]["required"], rows[0]["balance"]), (100, 100, 40))
        self.assertEqual((rows[1]["monthly_plan"], rows[1]["required"], rows[1]["balance"]), (150, 190, 10))
        self.assertEqual(rows[1]["actual_cumulative"], 240)

    def test_source_table_formula_example(self):
        data = default_data(2026)
        data["targets"][0]["plan"][:2] = [4085, 6790]
        data["events"] = [
            {"date": "10.01.2026", "target_ids": ["01"], "participants": 100},
            {"date": "10.02.2026", "target_ids": ["01"], "participants": 530},
        ]
        rows = calculate_target(data, "01")
        self.assertEqual(rows[0]["balance"], 3985)
        self.assertEqual(rows[1]["monthly_plan"], 2705)
        self.assertEqual(rows[1]["required"], 6690)
        self.assertEqual(rows[1]["balance"], 6160)

    def test_multiple_classifications(self):
        result = classify_event("Патриотическая лекция для волонтеров", "История и обучение", "Дом молодежи")
        self.assertTrue({"01", "02", "04", "05"}.issubset(result))

    def test_custom_target_classification(self):
        targets = [{"id": "new", "keywords": ["медиа", "журналист"]}]
        self.assertEqual(classify_event("Школа журналистики", "Медиапрактика", "", targets), ["new"])

    def test_old_data_is_migrated_to_editable_forms(self):
        data = normalize_data({
            "year": 2026,
            "targets": [{"id": "01", "name": "Старое название", "plan": [10]}],
            "events": [{"date": "10.01.2026", "name": "Событие", "target_ids": ["01"]}],
        })
        self.assertEqual(data["targets"][0]["code"], "1П")
        self.assertIn("report_settings", data)
        self.assertIn("card_labels", data)
        self.assertIsNone(data["events"][0]["participants"])
        self.assertEqual(data["events"][0]["press_release"], "")

    def test_draft_can_have_empty_post_event_fields(self):
        event = {
            "date": "10.01.2026", "time": "10:00", "place": "Дом молодежи",
            "name": "Запланированное мероприятие", "description": "", "target_ids": [],
            "participants": None, "link": "",
        }
        self.assertFalse(event_is_complete(event))

    def test_press_release_is_built_from_description(self):
        event = {
            "date": "10.01.2020", "time": "10:00", "place": "Дом молодежи", "name": "Открытая лекция",
            "description": "Участники обсудили историю города", "participants": 30, "link": "https://example.test",
        }
        text = build_press_release(event)
        self.assertIn("Открытая лекция", text)
        self.assertIn("Участники обсудили историю города", text)
        self.assertIn("30 человек", text)
        self.assertIn("https://example.test", text)

    def test_press_release_exports_to_word(self):
        data = default_data(2026)
        event = {
            "date": "10.01.2026", "time": "10:00", "place": "Дом молодежи", "name": "Открытая лекция",
            "description": "Участники обсудили историю города", "participants": 30, "link": "https://example.test",
        }
        text = build_press_release(event)
        with tempfile.TemporaryDirectory() as directory:
            path = export_press_release(data, event, text, Path(directory) / "press.docx")
            document = Document(path)
            content = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("ПРЕСС-РЕЛИЗ", content)
            self.assertIn("Участники обсудили историю города", content)
            self.assertIn("https://example.test", content)

    def test_cutoff_moves_actual_to_next_report_month(self):
        data = default_data(2026)
        data["events"] = [{"date": "25.01.2026", "target_ids": ["01"], "participants": 40}]
        rows = calculate_target(data, "01")
        self.assertEqual(rows[0]["actual"], 0)
        self.assertEqual(rows[1]["actual"], 40)

    def test_exports_open(self):
        data = default_data(2026)
        data["targets"][0]["plan"][0] = 100
        data["events"] = [{
            "id": "event-1", "date": "10.01.2026", "time": "10:00", "place": "Дом молодежи",
            "name": "Патриотическая лекция", "description": "Об истории России", "target_ids": ["01", "04"],
            "participants": 50, "link": "https://example.test",
        }, {
            "id": "event-2", "date": "10.02.2026", "time": "10:00", "place": "Дом молодежи",
            "name": "Февральское мероприятие", "description": "Не должно попасть в январский журнал", "target_ids": ["04"],
            "participants": 25, "link": "https://example.test/february",
        }]
        with tempfile.TemporaryDirectory() as directory:
            paths = export_all(data, Path(directory), 0)
            self.assertEqual(len(paths), 3)
            for path in paths:
                self.assertTrue(path.exists())
            appendix_1 = Document(paths[0])
            self.assertEqual([row.cells[1].text for row in appendix_1.tables[0].rows[2:]], ["01", "02", "03", "04", "05"])
            appendix_2 = Document(paths[1])
            appendix_2_text = " ".join(cell.text for table in appendix_2.tables for row in table.rows for cell in row.cells)
            self.assertIn("Патриотическая лекция", appendix_2_text)
            self.assertIn("10.01.2026", appendix_2_text)
            self.assertIn("50", appendix_2_text)
            self.assertIn("https://example.test", appendix_2_text)
            self.assertNotIn("Февральское мероприятие", appendix_2_text)
            workbook = load_workbook(paths[2], data_only=False)
            self.assertIn("Целевые показатели", workbook.sheetnames)
            self.assertIn("Январь", workbook.sheetnames)

    def test_appendix_2_uses_reporting_period_and_keeps_drafts(self):
        data = default_data(2026)
        data["events"] = [{
            "id": "draft", "date": "25.01.2026", "time": "10:00", "place": "Дом молодежи",
            "name": "Черновик февраля", "description": "", "target_ids": ["04"],
            "participants": None, "link": "",
        }]
        with tempfile.TemporaryDirectory() as directory:
            january_path = export_appendix_2(data, Path(directory) / "january.docx", 0)
            february_path = export_appendix_2(data, Path(directory) / "february.docx", 1)
            january_text = " ".join(cell.text for table in Document(january_path).tables for row in table.rows for cell in row.cells)
            february_text = " ".join(cell.text for table in Document(february_path).tables for row in table.rows for cell in row.cells)
            self.assertNotIn("Черновик февраля", january_text)
            self.assertIn("Черновик февраля", february_text)
            self.assertIn("Черновик февраля", preview_appendix_2(data, 1))

    def test_edited_report_form_and_new_target_are_exported(self):
        data = default_data(2027)
        data["report_settings"]["appendix_1_title"] = "НОВАЯ ФОРМА 2027"
        data["targets"].append({
            "id": "new-target", "code": "6М", "report_row": "06", "name": "Медиапроекты",
            "report_name": "Численность участников медиапроектов", "section": "МЕДИАПРОЕКТЫ",
            "keywords": ["медиа"], "plan": [0] * 12,
        })
        data["events"] = [{
            "id": "media", "date": "10.01.2027", "name": "Медиафорум", "time": "", "place": "",
            "description": "", "target_ids": ["new-target"], "participants": 12, "link": "",
        }]
        with tempfile.TemporaryDirectory() as directory:
            paths = export_all(data, Path(directory), 0)
            appendix_1 = Document(paths[0])
            self.assertIn("НОВАЯ ФОРМА 2027", " ".join(p.text for p in appendix_1.paragraphs))
            self.assertEqual(appendix_1.tables[0].rows[-1].cells[1].text, "06")
            appendix_2_text = " ".join(cell.text for table in Document(paths[1]).tables for row in table.rows for cell in row.cells)
            self.assertIn("МЕДИАПРОЕКТЫ", appendix_2_text)
            self.assertIn("Медиафорум", appendix_2_text)


if __name__ == "__main__":
    unittest.main()
