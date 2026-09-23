from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .core import (
    MONTHS,
    calculate_target,
    event_in_reporting_month,
    parse_event_date,
    reporting_period_bounds,
)


BLUE = "17324D"
PALE_BLUE = "DCE8F2"
GOLD = "D4A72C"
THIN = Side(style="thin", color="7B8794")


def _font(run, size=9, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def _shade(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), color)
    properties.append(shade)


def _set_cell(cell, text, bold=False, size=8, center=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(str(text))
    _font(run, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _setup_document(landscape=False):
    document = Document()
    section = document.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    return document


def _targets(data):
    def sort_key(item):
        row = str(item.get("report_row", ""))
        return (0, int(row)) if row.isdigit() else (1, row.lower())
    return sorted(data.get("targets", []), key=sort_key)


def export_appendix_1(data: dict, path: Path, through_month: int) -> Path:
    settings = data["report_settings"]
    document = _setup_document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run(settings["appendix_1_title"]), 12, True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(subtitle.add_run(
        f"{settings['appendix_1_subtitle']} за {MONTHS[through_month].lower()} {data['year']} года\n"
        f"{data['organization']}"
    ), 10, True)

    table = document.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = settings["appendix_1_columns"]
    for index, text in enumerate(headers):
        _set_cell(table.cell(0, index), text, True, 8, True)
        _shade(table.cell(0, index), PALE_BLUE)
    _set_cell(table.cell(1, 0), "1", False, 8, True)
    _set_cell(table.cell(1, 1), "2", False, 8, True)
    _set_cell(table.cell(1, 2), "3", False, 8, True)
    for target in _targets(data):
        row = table.add_row().cells
        calculated = calculate_target(data, target["id"])[through_month]
        _set_cell(row[0], target["report_name"], False, 8)
        _set_cell(row[1], target["report_row"], False, 8, True)
        _set_cell(row[2], calculated["actual_cumulative"], True, 9, True)
    widths = (Cm(14.7), Cm(1.5), Cm(2.7))
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = width
    document.save(path)
    return path


def export_appendix_2(data: dict, path: Path, through_month: int) -> Path:
    settings = data["report_settings"]
    document = _setup_document(landscape=True)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(title.add_run(settings["appendix_2_title"]), 12, True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    period_start, period_end = reporting_period_bounds(int(data["year"]), through_month)
    _font(subtitle.add_run(
        f"{data['organization']}\nза {MONTHS[through_month].lower()} {data['year']} года "
        f"({period_start:%d.%m.%Y}–{period_end:%d.%m.%Y})"
    ), 10, True)

    events = []
    for event in data.get("events", []):
        parsed = parse_event_date(event.get("date", ""))
        if parsed and event_in_reporting_month(event, int(data["year"]), through_month):
            events.append((parsed, event))
    events.sort(key=lambda item: (item[0], item[1].get("time", "")))

    for target in _targets(data):
        heading = document.add_table(rows=1, cols=1)
        heading.style = "Table Grid"
        heading.alignment = WD_TABLE_ALIGNMENT.CENTER
        _set_cell(heading.cell(0, 0), target["section"], True, 8, True)
        _shade(heading.cell(0, 0), PALE_BLUE)
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = settings["appendix_2_columns"]
        for index, text in enumerate(headers):
            _set_cell(table.cell(0, index), text, True, 7, True)
        selected = [(parsed, event) for parsed, event in events if target["id"] in event.get("target_ids", [])]
        for number, (parsed, event) in enumerate(selected, 1):
            row = table.add_row().cells
            participants = "" if event.get("participants") is None else event.get("participants")
            values = (number, parsed.strftime("%d.%m.%Y"), event.get("name", ""), participants, event.get("link", ""))
            for index, value in enumerate(values):
                _set_cell(row[index], value, False, 7, index in (0, 1, 3))
        if not selected:
            row = table.add_row().cells
            for cell in row:
                _set_cell(cell, "", False, 7)
        document.add_paragraph().paragraph_format.space_after = Pt(2)
    document.save(path)
    return path


def _style_sheet(sheet, widths):
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_monthly_summary(data: dict, path: Path) -> Path:
    workbook = Workbook()
    workbook.remove(workbook.active)
    target_map = {item["id"]: item for item in data["targets"]}
    labels = data["card_labels"]
    custom_fields = data.get("custom_fields", [])
    for month, month_name in enumerate(MONTHS):
        sheet = workbook.create_sheet(month_name)
        total_columns = 9 + len(custom_fields)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
        period_start, period_end = reporting_period_bounds(int(data["year"]), month)
        sheet["A1"] = (
            f"Учет мероприятий за {month_name.lower()} {data['year']} года "
            f"({period_start:%d.%m.%Y}–{period_end:%d.%m.%Y})"
        )
        sheet["A1"].font = Font(size=14, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor=BLUE)
        sheet["A1"].alignment = Alignment(horizontal="center")
        headers = (
            labels["date"], labels["time"], labels["place"], labels["name"], labels["description"],
            "Соответствие ЦП", labels["participants"], labels["link"], "Коды ЦП",
            *(field["label"] for field in custom_fields),
        )
        sheet.append(headers)
        events = []
        for event in data.get("events", []):
            parsed = parse_event_date(event.get("date", ""))
            if parsed and event_in_reporting_month(event, int(data["year"]), month):
                events.append((parsed, event))
        events.sort(key=lambda item: (item[0], item[1].get("time", "")))
        for parsed, event in events:
            target_names = ", ".join(target_map[x]["name"] for x in event.get("target_ids", []) if x in target_map)
            target_codes = ", ".join(target_map[x]["code"] for x in event.get("target_ids", []) if x in target_map)
            participants = "" if event.get("participants") is None else event.get("participants")
            custom_values = event.get("custom_fields", {})
            sheet.append((
                parsed.strftime("%d.%m.%Y"), event.get("time", ""), event.get("place", ""), event.get("name", ""),
                event.get("description", ""), target_names, participants, event.get("link", ""), target_codes,
                *(custom_values.get(field["id"], "") for field in custom_fields),
            ))
        _style_sheet(sheet, (13, 13, 30, 35, 55, 45, 16, 35, 14, *(25 for _field in custom_fields)))
        sheet.row_dimensions[1].height = 26
        sheet.row_dimensions[2].height = 38

    summary = workbook.create_sheet("Целевые показатели", 0)
    summary.merge_cells("A1:H1")
    summary["A1"] = f"Целевые показатели на {data['year']} год"
    summary["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=BLUE)
    summary["A1"].alignment = Alignment(horizontal="center")
    summary.append(("Направление", "Месяц", "План нарастающим", "Разница между нарастающими", "Реальный охват в месяц", "Реальный охват нарастающим", "Разница: необходимо - сделано", "Необходимый охват с остатком"))
    custom_names = {item["id"]: item["name"] for item in data["targets"]}
    for target in data["targets"]:
        for row in calculate_target(data, target["id"]):
            summary.append((custom_names[target["id"]], MONTHS[row["month"]], row["cumulative_plan"], row["monthly_plan"], row["actual"], row["actual_cumulative"], row["balance"], row["required"]))
    _style_sheet(summary, (48, 14, 20, 24, 21, 24, 25, 27))
    summary.freeze_panes = "C3"
    workbook.save(path)
    return path


def export_all(data: dict, folder: Path, through_month: int) -> list[Path]:
    folder.mkdir(parents=True, exist_ok=True)
    suffix = f"{data['year']}_{through_month + 1:02d}"
    return [
        export_appendix_1(data, folder / f"Приложение_1_{suffix}.docx", through_month),
        export_appendix_2(data, folder / f"Приложение_2_{suffix}.docx", through_month),
        export_monthly_summary(data, folder / f"Сводные_данные_{data['year']}.xlsx"),
    ]


def export_press_release(data: dict, event: dict, text: str, path: Path) -> Path:
    document = _setup_document()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(heading.add_run("ПРЕСС-РЕЛИЗ"), 14, True)
    organization = document.add_paragraph()
    organization.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _font(organization.add_run(data.get("organization", "")), 9)
    document.add_paragraph()
    for index, paragraph_text in enumerate(part.strip() for part in text.split("\n\n") if part.strip()):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
        _font(paragraph.add_run(paragraph_text), 12 if index == 0 else 11, index == 0)
        paragraph.paragraph_format.space_after = Pt(8)
        if index:
            paragraph.paragraph_format.first_line_indent = Cm(1.25)
    details = document.add_paragraph()
    details.paragraph_format.space_before = Pt(12)
    _font(details.add_run(f"Дата мероприятия: {event.get('date', '')}\n"), 9)
    _font(details.add_run(f"Ссылка: {event.get('link', '')}"), 9)
    document.save(path)
    return path


def preview_appendix_1(data: dict, through_month: int) -> str:
    settings = data["report_settings"]
    lines = [
        settings["appendix_1_title"],
        f"{settings['appendix_1_subtitle']} за {MONTHS[through_month].lower()} {data['year']} года, нарастающим итогом",
        data["organization"],
        "",
    ]
    for target in _targets(data):
        value = calculate_target(data, target["id"])[through_month]["actual_cumulative"]
        lines.extend((f"Строка {target['report_row']}: {value}", target["report_name"], ""))
    return "\n".join(lines)


def preview_appendix_2(data: dict, month: int) -> str:
    settings = data["report_settings"]
    year = int(data["year"])
    period_start, period_end = reporting_period_bounds(year, month)
    events = []
    for event in data.get("events", []):
        parsed = parse_event_date(event.get("date", ""))
        if parsed and event_in_reporting_month(event, year, month):
            events.append((parsed, event))
    events.sort(key=lambda item: (item[0], item[1].get("time", "")))
    lines = [
        settings["appendix_2_title"],
        f"Отчетный период: {period_start:%d.%m.%Y}–{period_end:%d.%m.%Y}",
        data["organization"],
        f"Всего карточек за период: {len(events)}",
        "",
    ]
    for target in _targets(data):
        selected = [(parsed, event) for parsed, event in events if target["id"] in event.get("target_ids", [])]
        lines.append(target["section"])
        if not selected:
            lines.append("Нет мероприятий")
        for number, (parsed, event) in enumerate(selected, 1):
            participants = "не заполнено" if event.get("participants") is None else str(event["participants"])
            link = event.get("link") or "не заполнено"
            lines.append(f"{number}. {parsed:%d.%m.%Y} | {event.get('name', '')} | {participants} | {link}")
        lines.append("")
    unassigned = [event for _parsed, event in events if not event.get("target_ids")]
    if unassigned:
        lines.append("ВНИМАНИЕ: НЕ РАСПРЕДЕЛЕНЫ ПО ЦЕЛЕВЫМ ПОКАЗАТЕЛЯМ")
        lines.extend(f"• {event.get('date', '')} — {event.get('name', '')}" for event in unassigned)
    return "\n".join(lines)
